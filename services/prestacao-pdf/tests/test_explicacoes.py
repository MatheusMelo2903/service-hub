# -*- coding: utf-8 -*-
"""Explicação dos números que assustam na prestação v13 (déficit, caixa caindo, mês negativo).

Testes de LÓGICA usam um condomínio fictício com números redondos e rodam em qualquer
clone. O teste com dados reais só roda quando PRESTACAO_DADOS_REAIS aponta para um JSON
{"dados":..., "rascunho":...} (não vai para o git); sem ele, pula.

NÃO há constantes com valores financeiros reais neste arquivo.
"""
from __future__ import annotations

import copy
import io
import json
import os
import re
import sys
import zipfile

AQUI = os.path.dirname(__file__)
sys.path.insert(0, os.path.join(AQUI, ".."))
sys.path.insert(0, os.path.join(AQUI, "..", "vendor", "skill_v13", "scripts"))

import pytest

from app import justificativas as J
from app import v13
import montar as M
import gerador as G
from auditar_apresentacao import auditar

HIFEN = re.compile(r"\s[-‐-―−]\s|[–—]")


def _dados_ficticios(com_alerta: bool = True) -> dict:
    """Seis meses de um condomínio fictício. Com alerta: reforma da piscina em mar e abr
    paga com o fundo de reserva (déficit no período, dois meses negativos)."""
    per = [f"2026-{m:02d}" for m in range(1, 7)]
    lab = ["Jan/26", "Fev/26", "Mar/26", "Abr/26", "Mai/26", "Jun/26"]
    reforma = [0, 0, 60000.0, 40000.0, 0, 0] if com_alerta else [0.0] * 6
    ferias = [0, 0, 0, 0, 8000.0, 0] if com_alerta else [0.0] * 6
    rec = [50000.0] * 6
    rubs_man = [{"nome": "Contrato Elevador", "serie": [5000.0] * 6, "total": 30000.0, "itens": [], "status": "sem_detalhe"}]
    if com_alerta:
        rubs_man.append({"nome": "Reforma da Piscina", "serie": reforma, "total": 100000.0, "itens": [], "status": "sem_detalhe",
                         "itens_mes": {"2026-03": [["Ref. a primeira medição - reforma piscina", 60000.0]],
                                       "2026-04": [["Referente a segunda medição da reforma", 40000.0]]}})
    rubs_pes = [{"nome": "Salários", "serie": [30000.0] * 6, "total": 180000.0, "itens": [], "status": "sem_detalhe"}]
    if com_alerta:
        rubs_pes.append({"nome": "Férias", "serie": ferias, "total": 8000.0, "itens": [], "status": "sem_detalhe",
                         "itens_mes": {"2026-05": [["Ref. a Fulano de Tal (01/05/2026 a 30/05/2026", 8000.0]]}})
    cats = [{"nome": "DESPESA COM PESSOAL", "rubricas": rubs_pes}, {"nome": "MANUTENÇÃO", "rubricas": rubs_man}]
    for c in cats:
        c["total"] = round(sum(r["total"] for r in c["rubricas"]), 2)
    desp = [round(sum(r["serie"][i] for c in cats for r in c["rubricas"]), 2) for i in range(6)]
    ini = 200000.0
    saldos, s = [], ini
    for r_, d_ in zip(rec, desp):
        s = round(s + r_ - d_, 2); saldos.append(s)
    if com_alerta:
        fundo = [150000.0, 150000.0, 150000.0, 100000.0, 60000.0, 60000.0, 60000.0]
    else:
        fundo = [150000.0] * 7
    tot = [ini] + saldos
    corrente = [round(t - f, 2) for t, f in zip(tot, fundo)]
    contas = [{"nome": "Conta Corrente", "ini": corrente[0], "fim": corrente[-1], "saldos": corrente},
              {"nome": "Fundo de Reserva", "ini": fundo[0], "fim": fundo[-1], "saldos": fundo},
              {"nome": "Poupança Antiga", "ini": 0.0, "fim": 0.0, "saldos": [0.0] * 7}]
    return dict(condominio="Condomínio Exemplo Fictício", modo="detalhado", periodo=per, meses_label=lab, contas=contas,
                saldo_anterior=ini, saldo_final=saldos[-1], receitas_mes=rec, despesas_mes=desp, saldo_fim_mes=saldos,
                receitas=[{"nome": "Taxa Condominial", "serie": rec, "total": sum(rec)}], categorias=cats)


CURADORIA = {
    "resumo": "O déficit veio da reforma da piscina (R$ 100 mil, em março e abril), paga com o fundo de reserva — aprovada em assembleia.",
    "caixa": "Caiu pela reforma da piscina, paga com o fundo de reserva",
    "meses": {"mar/26": "Primeira medição da reforma da piscina (R$ 60 mil)",
              "abr/26": "Segunda medição da reforma da piscina (R$ 40 mil)",
              "dez/30": "mês que não existe"},
    "rubricas": {"MANUTENÇÃO|Reforma da Piscina": "Reforma da piscina aprovada em assembleia - paga com o fundo de reserva.",
                 "MANUTENÇÃO|Rubrica Inexistente": "não entra"},
    "contas": "O fundo de reserva pagou a reforma: saíram R$ 90 mil em março e abril.",
    "chave_estranha": "não entra",
}


def _comp(dados, exp=None):
    comp_ia = {"rubricas": {}, "categorias": {}, "receitas_nomes": {}}
    if exp is not None:
        comp_ia["explicacoes"] = exp
    return v13.completar_composicao(comp_ia, dados, {}, "Condomínio Exemplo Fictício")


def _cfg(dados, comp):
    """O que o v13.montar entrega ao montar.py: composição com as explicações já juntadas."""
    comp = dict(comp)
    exp = v13.juntar_explicacoes(dados, comp)
    if exp:
        comp["explicacoes"] = exp
    else:
        comp.pop("explicacoes", None)
    return M.montar_cfg(json.loads(json.dumps(dados)), comp)


def _pptx(cfg, tmp_path, nome="saida.pptx"):
    p = str(tmp_path / nome)
    G.montar(copy.deepcopy(cfg), p)
    return p


def _textos(pptx_path):
    from pptx import Presentation
    prs = Presentation(pptx_path)
    return [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).replace("\u00a0", " ") for s in prs.slides]


# ---------------- saneamento ----------------

def test_completar_composicao_saneia_explicacoes():
    dados = _dados_ficticios()
    comp = _comp(dados, CURADORIA)
    e = comp["explicacoes"]
    assert set(e) <= {"resumo", "caixa", "meses", "rubricas", "contas"}
    assert "dez/30" not in e["meses"] and set(e["meses"]) == {"mar/26", "abr/26"}
    assert list(e["rubricas"]) == ["MANUTENÇÃO|Reforma da Piscina"]
    for t in [e["resumo"], e["caixa"], e["contas"], *e["meses"].values(), *e["rubricas"].values()]:
        assert not HIFEN.search(t), t
    assert len(e["resumo"]) <= 240 and len(e["caixa"]) <= 90 and len(e["contas"]) <= 220
    # Guardada nos dados: vale também quando o main.py cai na composição automática.
    assert dados[v13._EXP_GUARDADA] == e
    longo = dict(CURADORIA, resumo="Frase curta de abertura. " + "palavra " * 60)
    assert len(_comp(dados, longo)["explicacoes"]["resumo"]) <= 240


def test_explicacao_da_curadoria_sobrevive_ao_fallback_automatico():
    dados = _dados_ficticios()
    _comp(dados, CURADORIA)                                # IA reprovou depois disso
    auto = v13.composicao_automatica(dados, {}, "Condomínio Exemplo Fictício")
    exp = v13.juntar_explicacoes(dados, auto)
    assert exp["caixa"] == "Caiu pela reforma da piscina, paga com o fundo de reserva"


def test_texto_padrao_completa_o_que_a_curadoria_nao_escreveu():
    dados = _dados_ficticios()
    exp = v13.juntar_explicacoes(dados, _comp(dados, {"caixa": "Texto da curadoria para o caixa"}))
    assert exp["caixa"] == "Texto da curadoria para o caixa"
    assert exp["resumo"].startswith("O déficit veio de") and "Fundo de reserva" in exp["resumo"]
    assert "MANUTENÇÃO|Reforma da Piscina" in exp["rubricas"]


def test_texto_padrao_limpa_rotulos():
    s = J.analisar(_dados_ficticios())
    tudo = json.dumps(s["textos"], ensure_ascii=False)
    assert "Fulano" not in tudo and "Ref." not in tudo and "Referente" not in tudo
    assert "férias da equipe (01/05/2026 a 30/05/2026)" in tudo
    assert not HIFEN.search(tudo)
    assert J.limpa_rotulo("Ref. a quinta medição - reforma") == "quinta medição, reforma"
    assert J.sem_hifen("Eletrodoméstico-Eletroeletrônico") == "Eletrodoméstico e Eletroeletrônico"


# ---------------- sem alerta: nada muda ----------------

def test_sem_alerta_cfg_e_slides_identicos(tmp_path):
    dados = _dados_ficticios(com_alerta=False)
    assert not J.analisar(dados)["ha_alerta"]
    base = M.montar_cfg(json.loads(json.dumps(dados)), _comp(dados))
    com_exp = _cfg(dados, _comp(dados, CURADORIA))          # curadoria mandou texto, mas não há alerta
    assert com_exp == base
    assert not {"explicacoes", "contas", "rubricas_marcadas"} & set(base)
    a, b = _pptx(base, tmp_path, "a.pptx"), _pptx(com_exp, tmp_path, "b.pptx")
    za, zb = zipfile.ZipFile(a), zipfile.ZipFile(b)
    nomes = sorted(n for n in za.namelist() if n.startswith(("ppt/slides/", "ppt/charts/")))
    assert nomes == sorted(n for n in zb.namelist() if n.startswith(("ppt/slides/", "ppt/charts/")))
    assert all(za.read(n) == zb.read(n) for n in nomes)


# ---------------- com alerta ----------------

def _confere_deck(dados, comp, tmp_path, frases):
    cfg = _cfg(dados, comp)
    assert cfg["explicacoes"] and cfg["contas"]
    contas = cfg["contas"]
    assert abs(sum(c["ini"] for c in contas) - cfg["saldo_anterior"]) < 0.02
    assert abs(sum(c["fim"] for c in contas) - cfg["saldo_final"]) < 0.02
    p = _pptx(cfg, tmp_path)
    tx = _textos(p)
    problemas, _ = auditar(p)
    assert problemas == [], problemas
    assert "Onde está o dinheiro" in tx[4]                    # logo depois do patrimônio
    assert G.brl(cfg["saldo_anterior"]) in tx[4] and G.brl(cfg["saldo_final"]) in tx[4]
    for onde, frase in frases:
        assert frase in tx[onde], (onde, frase)
    return cfg, tx


def test_com_alerta_explicacao_em_todos_os_slides(tmp_path):
    dados = _dados_ficticios()
    comp = _comp(dados, CURADORIA)
    cfg, tx = _confere_deck(dados, comp, tmp_path, [
        (0, "VARIAÇÃO NO PERÍODO"), (0, CURADORIA["caixa"]),
        (1, "POR QUE O RESULTADO FICOU NEGATIVO"), (1, "Caiu pela reforma da piscina"),
        (3, "DE ONDE SAIU O DINHEIRO"), (3, "Fundo de Reserva de R$ 150 mil para R$ 60 mil"),
        (5, "Primeira medição da reforma da piscina"),
    ])
    assert "Poupança Antiga" not in tx[4]                       # conta zerada nos dois extremos fica fora
    det = [t for t in tx if "Reforma da Piscina" in t and "TOTAL DA CATEGORIA" in t]
    assert det and "▸ Reforma da piscina aprovada em assembleia, paga com o fundo de reserva." in det[0]
    # A composição (com R$) segue intacta: o destaque fica em campo separado.
    rub = [r for c in cfg["despesas"] for r in c["rubricas"] if r["chave"] == "Reforma da Piscina"][0]
    assert rub["destaque"] and "R$" not in rub["destaque"]


def test_com_alerta_so_texto_padrao(tmp_path):
    dados = _dados_ficticios()
    _confere_deck(dados, _comp(dados), tmp_path, [
        (0, "Uso planejado de reserva"), (1, "O déficit veio de"), (5, "MESES NEGATIVOS"),
    ])


def test_contas_que_nao_fecham_nao_viram_slide(tmp_path):
    dados = _dados_ficticios()
    dados["contas"][0]["fim"] += 5.0
    cfg = _cfg(dados, _comp(dados, CURADORIA))
    assert cfg["explicacoes"] and "contas" not in cfg
    tx = _textos(_pptx(cfg, tmp_path))
    assert not any("Onde está o dinheiro" in t for t in tx)


# ---------------- dados reais (opcional) ----------------

REAIS = os.environ.get("PRESTACAO_DADOS_REAIS")


@pytest.mark.skipif(not (REAIS and os.path.exists(REAIS)), reason="sem dados reais (PRESTACAO_DADOS_REAIS)")
def test_dados_reais_com_deficit(tmp_path):
    raw = json.load(open(REAIS, encoding="utf-8"))
    dados = raw["dados"]
    comp = v13.completar_composicao({"rubricas": {}, "categorias": {}, "receitas_nomes": {}}, dados, raw.get("rascunho") or {})
    sinais = J.analisar(dados)
    if not sinais["ha_alerta"]:
        pytest.skip("período sem alerta")
    cfg = _cfg(dados, comp)
    e = cfg["explicacoes"]
    frases = [(0, e["caixa"])] if e.get("caixa") else []
    _confere_deck(dados, comp, tmp_path, frases + [(1, e["resumo"][:30])])


# ---------------- rótulos da composição (rascunho, automática e IA) ----------------

RASC_FERIAS = ("Ref. a Fulano (15/06/2026 a 04/07/2026 (2 lançamentos) R$ 5.000,00 · Ref. Férias Beltrana 08/09 a 07/10 R$ 2.000,00"
               " · Ref. à Ciclana Souza (16 à 30/12/2025 R$ 1.000,00")


def test_limpa_composicao_tira_nome_e_mantem_valores():
    t = J.limpa_composicao(RASC_FERIAS, "Férias")
    assert t == ("férias da equipe (15/06/2026 a 04/07/2026) (2 lançamentos) R$ 5.000,00 · férias da equipe (08/09 a 07/10) "
                 "R$ 2.000,00 · férias da equipe (16 a 30/12/2025) R$ 1.000,00")
    assert J.limpa_composicao("Ref. a quinta medição - reforma R$ 10,00 · lançamentos sem descrição (3 lançamentos) R$ 5,00", "Obra") \
        == "quinta medição, reforma R$ 10,00 · lançamentos sem descrição (3 lançamentos) R$ 5,00"
    # Sem nome de pessoa, o rótulo de pessoal fica (só sem o "Ref.").
    assert J.limpa_composicao("Ref. a 2ª parcela R$ 300,00", "13 Salário") == "2ª parcela R$ 300,00"
    assert J.limpa_composicao("Lançamentos do período ({n} meses)") == "Lançamentos do período ({n} meses)"


def test_composicao_automatica_e_ia_sem_nome_de_funcionario(tmp_path):
    dados = _dados_ficticios()
    ferias = [r for c in dados["categorias"] for r in c["rubricas"] if r["nome"] == "Férias"][0]
    ferias["status"] = "exato"
    rasc = {"Férias": {"texto": "Ref. a Fulano de Tal (01/05/2026 a 30/05/2026 R$ 8.000,00", "fecha": True, "status": "exato"}}
    auto = v13.composicao_automatica(dados, rasc)
    assert auto["rubricas"]["Férias"] == {"texto": "férias da equipe (01/05/2026 a 30/05/2026) R$ 8.000,00", "fecha": True}
    assert v13.limpar_rascunho(rasc)["Férias"]["texto"] == auto["rubricas"]["Férias"]["texto"]
    comp = v13.completar_composicao({"rubricas": {"Férias": {"texto": "Férias de Fulano de Tal R$ 8.000,00", "fecha": True}}}, dados, rasc)
    assert comp["rubricas"]["Férias"] == {"texto": "férias da equipe R$ 8.000,00", "fecha": True}
    # A conferência de soma ("fecha") continua passando no gerador.
    cfg = _cfg(dados, comp)
    p = _pptx(cfg, tmp_path)
    assert "Fulano" not in " ".join(_textos(p))
