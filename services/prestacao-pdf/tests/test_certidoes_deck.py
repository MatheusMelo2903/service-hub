# -*- coding: utf-8 -*-
"""Certidões dentro da prestação: um slide por página de certidão, no fim do PowerPoint,
com o rodapé copiado da apresentação e a numeração continuando."""
import io
import os
import tempfile

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app import certidoes_deck as C


def _pdf_de_teste(paginas=1):
    imgs = []
    for p in range(paginas):
        im = Image.new('RGB', (1240, 1754), 'white')  # A4 a 150 dpi
        d = ImageDraw.Draw(im)
        for i in range(40):
            d.text((120, 150 + i * 30), f'Linha {i} da pagina {p + 1}: nada consta.', fill='black')
        imgs.append(im)
    buf = io.BytesIO()
    imgs[0].save(buf, 'PDF', save_all=True, append_images=imgs[1:])
    return buf.getvalue()


def _deck_de_teste():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for _ in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(9), Inches(0.3))
        tb.text_frame.text = 'CONDOMÍNIO TESTE    •    EXERCÍCIO OUT/2025 A SET/2026'
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_slides_no_fim_do_pptx():
    certs = [
        {'tipo': 'fgts', 'nome': 'FGTS (Caixa)', 'situacao': 'negativa', 'emitida_em': '2026-10-23', 'validade': '2026-11-21', 'pdf': _pdf_de_teste(1)},
        {'tipo': 'federal', 'nome': 'Federal (Receita e PGFN)', 'situacao': 'negativa', 'emitida_em': '2026-06-24', 'validade': '2026-12-21', 'pdf': _pdf_de_teste(2)},
    ]
    with tempfile.TemporaryDirectory() as tmp:
        r = C.anexar(_deck_de_teste(), None, certs, tmp)
    assert r['pdf'] is None
    assert r['slides'] == 3  # federal com 2 páginas + FGTS
    prs = Presentation(io.BytesIO(r['pptx']))
    assert len(prs.slides) == 6
    textos = [' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    # ordem fixa: federal antes do FGTS
    assert 'Federal' in textos[3] and 'PÁGINA 1 DE 2' in textos[3]
    assert 'PÁGINA 2 DE 2' in textos[4]
    assert 'FGTS' in textos[5] and 'REGULAR' in textos[5]
    assert 'EXERCÍCIO OUT/2025 A SET/2026' in textos[5]
    assert textos[5].strip().endswith('6')  # numeração continua


def test_sem_certidao_recusa():
    with tempfile.TemporaryDirectory() as tmp:
        try:
            C.anexar(_deck_de_teste(), None, [], tmp)
        except C.ErroPrestacao:
            return
    raise AssertionError('deveria recusar sem certidão')
