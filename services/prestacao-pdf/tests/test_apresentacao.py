# -*- coding: utf-8 -*-
"""Motor de layouts da apresentação de assembleia: um roteiro com todos os tipos de
slide, fotos geradas com Pillow e um PDF gerado na hora. Confere PPTX e PDF válidos,
número de slides, avisos (tipo desconhecido, foto inexistente, avisos do roteiro),
paginação de tabela e texto longos, logo do cliente e a rota HTTP."""
import base64
import io
import os
import re
import subprocess
import tempfile

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation

from app import apresentacao as A


def _foto(cor=(60, 140, 90), w=1200, h=800):
    im = Image.new('RGB', (w, h), cor)
    ImageDraw.Draw(im).ellipse([w // 4, h // 4, w // 2, h // 2], fill=(230, 200, 60))
    buf = io.BytesIO()
    im.save(buf, 'JPEG')
    return buf.getvalue()


def _png_transparente():
    im = Image.new('RGBA', (600, 400), (0, 0, 0, 0))
    ImageDraw.Draw(im).rectangle([100, 100, 500, 300], fill=(20, 20, 20, 255))
    buf = io.BytesIO()
    im.save(buf, 'PNG')
    return buf.getvalue()


def _pdf(paginas=2):
    imgs = []
    for p in range(paginas):
        im = Image.new('RGB', (1240, 1754), 'white')
        d = ImageDraw.Draw(im)
        for i in range(30):
            d.text((120, 150 + i * 30), f'Linha {i} da pagina {p + 1}', fill='black')
        imgs.append(im)
    buf = io.BytesIO()
    imgs[0].save(buf, 'PDF', save_all=True, append_images=imgs[1:])
    return buf.getvalue()


def _logo():
    im = Image.new('RGB', (800, 300), 'white')
    d = ImageDraw.Draw(im)
    d.ellipse([20, 20, 280, 280], fill=(20, 110, 70))
    d.rectangle([320, 100, 760, 200], fill=(200, 150, 30))
    buf = io.BytesIO()
    im.save(buf, 'PNG')
    return buf.getvalue()


LONGO = ('O condomínio recebeu a notificação da prefeitura sobre a calçada. ' * 30).strip()

ROTEIRO = {
    'titulo_arquivo': 'AGE teste',
    'avisos': ['A análise da diretoria diverge do demonstrativo em setembro.'],
    'slides': [
        {'tipo': 'capa', 'assembleia': 'Assembleia Geral Extraordinária', 'sigla': 'AGE',
         'condominio': 'Condomínio Teste Clube', 'data': '16 de outubro de 2024',
         'chamadas': ['1ª chamada às 19h', '2ª chamada às 19h30'], 'local': 'Salão de festas'},
        {'tipo': 'orientacoes', 'blocos': [{'titulo': 'Quórum', 'topicos': ['Metade dos condôminos']},
                                           {'titulo': 'Procuração', 'topicos': ['Uma por condômino']}]},
        {'tipo': 'pauta', 'itens': [{'titulo': 'Reforma do parquinho', 'etiqueta': 'VOTAÇÃO'},
                                    'Pintura das garagens', 'Regimento interno', 'Eleição do conselho']},
        {'tipo': 'item', 'numero': 1, 'titulo': 'Reforma do parquinho', 'subtitulo': 'Piso e brinquedos',
         'etiqueta': 'VOTAÇÃO', 'apresentador': 'Síndica', 'fonte': ['edital, página 1']},
        {'tipo': 'texto', 'numero': 1, 'titulo': 'Contexto', 'paragrafos': ['Texto curto.'], 'destaque': 'Frase em destaque.'},
        {'tipo': 'texto', 'numero': 1, 'titulo': 'Texto longo', 'paragrafos': [LONGO, LONGO]},
        {'tipo': 'topicos', 'numero': 1, 'titulo': 'Situação', 'topicos': ['Um', 'Dois', 'Três']},
        {'tipo': 'colunas', 'numero': 1, 'titulo': 'Opções', 'colunas': [{'nome': 'A', 'topicos': ['x']}, {'nome': 'B', 'topicos': ['y']}]},
        {'tipo': 'orcamento', 'numero': 1, 'titulo': 'Orçamento', 'fornecedor': 'Fornecedor A',
         'itens': [{'descricao': 'Piso', 'valor': 1000.0}, {'descricao': 'Mesa', 'valor': 500.0}], 'total': 1500.0,
         'pagamento': 'À vista', 'prazo': '30 dias'},
        {'tipo': 'orcamento', 'numero': 1, 'titulo': 'Orçamento B', 'fornecedor': 'Fornecedor B',
         'mao_de_obra': 34605.0, 'materiais': 9292.45, 'total': 43897.45, 'documento': 'd1', 'pagina': 1},
        {'tipo': 'comparativo', 'numero': 1, 'titulo': 'Comparativo', 'mostrar_media': True,
         'opcoes': [{'nome': 'A', 'valor': 1500.0}, {'nome': 'B', 'valor': 1800.0}, {'nome': 'C', 'valor': 1700.0},
                    {'nome': 'D', 'valor': 2000.0}, {'nome': 'E', 'valor': 1900.0}]},
        {'tipo': 'tabela', 'numero': 1, 'titulo': 'Tabela longa', 'colunas': ['Serviço', 'Valor'],
         'linhas': [[f'Serviço {i}', f'R$ {i}.000,00'] for i in range(1, 26)], 'total': ['Total', 'R$ 325.000,00']},
        {'tipo': 'numeros', 'numero': 1, 'titulo': 'Saldos', 'cards': [{'rotulo': 'Conta', 'valor': 'R$ 1.000,00'}, {'rotulo': 'Fundo', 'valor': 2500.5}]},
        {'tipo': 'barras', 'numero': 1, 'titulo': 'Despesas', 'mostrar_percentual': True,
         'itens': [{'nome': 'Pessoal', 'valor': 800.0}, {'nome': 'Água', 'valor': 200.0}, {'nome': 'Estorno', 'valor': -50.0}]},
        {'tipo': 'cards', 'numero': 1, 'titulo': 'Frentes', 'cards': [{'titulo': 'Queijaria', 'texto': 'Ajustes', 'status': 'EM ADEQUAÇÃO'},
                                                                     {'titulo': 'Cestas', 'texto': 'Entregues', 'status': 'ENTREGUE'}]},
        {'tipo': 'passos', 'numero': 1, 'titulo': 'Etapas', 'passos': [{'titulo': 'Limpeza', 'texto': 'Retirada'}, {'titulo': 'Obra', 'texto': 'Execução'}]},
        {'tipo': 'pontos', 'numero': 1, 'titulo': 'Pontos', 'pontos': [{'texto': 'Cercamento', 'origem': 'Item 2'}]},
        {'tipo': 'fotos', 'numero': 1, 'titulo': 'Fotos', 'fotos': [{'id': 'f1', 'legenda': 'Uma'}, {'id': 'f2'}, {'id': 'f9'}]},
        {'tipo': 'foto_texto', 'numero': 1, 'titulo': 'Brinquedo', 'foto': 'f3', 'topicos': ['Balanço'], 'valor': 'R$ 4.615,60'},
        {'tipo': 'antes_depois', 'numero': 1, 'titulo': 'Salão', 'antes': 'f1', 'depois': 'f2', 'topicos': ['Piso novo']},
        {'tipo': 'documento', 'numero': 1, 'titulo': 'Projeto', 'documento': 'd1', 'paginas': [1, 2, 7]},
        {'tipo': 'regimento', 'numero': 2, 'titulo': 'Capítulo I', 'artigo': 'Art. 1', 'texto_atual': None, 'proposta': 'Novo texto.'},
        {'tipo': 'custeio', 'numero': 2, 'titulo': 'Custeio', 'valor_total': 21000.0, 'unidades': 440,
         'opcoes': [{'nome': 'Opção 1', 'parcelas': 1, 'valor_parcela': 47.73}]},
        {'tipo': 'votacao', 'numero': 2, 'titulo': 'Votação', 'pergunta': 'Aprovar?', 'regra': 'Maioria simples.'},
        {'tipo': 'eleicao', 'numero': 3, 'titulo': 'Eleição', 'cargos': [{'cargo': 'Síndico', 'mandato': '2025 a 2026'}]},
        {'tipo': 'linha_do_tempo', 'numero': 3, 'titulo': 'Processo', 'eventos': [{'data': '01/01/2024', 'titulo': 'Início'},
                                                                                 {'data': '02/02/2024', 'titulo': 'Fim', 'texto': 'Acordo.'}]},
        {'tipo': 'desconhecido', 'titulo': 'teste'},
        {'tipo': 'encerramento'},
    ],
}
FOTOS = {'f1': _foto(), 'f2': _foto((150, 80, 60), 800, 1200), 'f3': _png_transparente()}
DOCS = {'d1': _pdf(2)}


@pytest.fixture(scope='module')
def gerado():
    with tempfile.TemporaryDirectory() as tmp:
        pptx, pdf, avisos, n = A.montar(ROTEIRO, FOTOS, DOCS, tmp, _logo())
    return pptx, pdf, avisos, n


def _textos(pptx):
    prs = Presentation(io.BytesIO(pptx))
    # R$ e o valor vão com espaço não separável no slide; aqui volta a espaço comum.
    return prs, [' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).replace('\u00a0', ' ') for s in prs.slides]


def test_pptx_e_pdf_validos(gerado):
    pptx, pdf, _, n = gerado
    assert pdf[:4] == b'%PDF'
    prs, _ = _textos(pptx)
    assert len(prs.slides) == n
    with tempfile.NamedTemporaryFile(suffix='.pdf') as f:
        f.write(pdf)
        f.flush()
        info = subprocess.run(['pdfinfo', f.name], capture_output=True, text=True).stdout
    assert int(re.search(r'Pages:\s+(\d+)', info).group(1)) == n


def test_todos_os_tipos_desenhados(gerado):
    pptx, _, avisos, n = gerado
    _, textos = _textos(pptx)
    tudo = '\n'.join(textos)
    for trecho in ('AGE', 'ORIENTAÇÕES DO EDITAL', 'Pauta da', 'ITEM 1 DE 4', 'Frase em destaque', 'MÃO DE OBRA',
                   'MENOR VALOR', 'MÉDIA DOS 5 ORÇAMENTOS', 'R$ 2.500,50', 'EM ADEQUAÇÃO', 'PONTO 1', 'ANTES', 'DEPOIS',
                   'PÁGINA DO DOCUMENTO', 'Texto inexistente hoje', 'POR UNIDADE', 'R$ 47,73', 'O QUE SERÁ VOTADO',
                   'REGRA DE VOTAÇÃO', 'CARGO EM ELEIÇÃO', '01/01/2024', 'Agradecemos a presença de todos!'):
        assert trecho in tudo, trecho
    # rodapé com condomínio e data e numeração em todos os slides de conteúdo
    assert 'Grupo Service  ·  Condomínio Teste Clube  ·  16 de outubro de 2024' in textos[1]
    assert f'16 de outubro de 2024 {n} ' in textos[-1]   # numeração no rodapé do último slide
    # nada de "fonte" do roteiro no slide
    assert 'edital, página 1' not in tudo


def test_avisos(gerado):
    _, _, avisos, _ = gerado
    junto = '\n'.join(avisos)
    assert avisos[0].startswith('Roteiro: A análise da diretoria')
    assert 'tipo "desconhecido" não existe' in junto
    assert 'foto "f9" não foi enviada' in junto
    assert 'página 7 não existe' in junto


def test_paginacao_tabela_e_texto(gerado):
    pptx, _, _, _ = gerado
    _, textos = _textos(pptx)
    tab = [t for t in textos if 'Tabela longa' in t]
    assert len(tab) >= 3                       # 25 linhas, até 12 por slide
    assert all('(continuação)' in t for t in tab[1:])
    assert 'Serviço 25' in tab[-1] and 'TOTAL' in tab[-1]
    txt = [t for t in textos if 'Texto longo' in t]
    assert len(txt) >= 2 and '(continuação)' in txt[1]


def test_sem_hifen_de_pontuacao():
    with tempfile.TemporaryDirectory() as tmp:
        r = {'slides': [{'tipo': 'capa', 'condominio': 'Cond X', 'data': '01/01/2025'},
                        {'tipo': 'topicos', 'titulo': 'Itens - gerais', 'topicos': ['Portão — quebrado', 'CNPJ 12.345.678/0001-90']}]}
        prs, _ = A.desenhar(r, {}, {}, tmp)
    tudo = ' '.join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert ' - ' not in tudo and '—' not in tudo
    assert 'Portão, quebrado' in tudo and '0001-90' in tudo


def test_roteiro_ruim_nao_quebra():
    with tempfile.TemporaryDirectory() as tmp:
        r = {'slides': [None, {'tipo': 'orcamento'}, {'tipo': 'comparativo', 'opcoes': 'x'}, {'tipo': 'tabela', 'linhas': [[1, None]]},
                        {'tipo': 'custeio', 'valor_total': 'abc'}, {'tipo': 'fotos', 'fotos': [{'id': 'nada'}]}, {'tipo': 'encerramento'}]}
        prs, avisos = A.desenhar(r, {}, {}, tmp)
    assert len(prs.slides) >= 2
    assert any('não tem capa' in a for a in avisos)
    with tempfile.TemporaryDirectory() as tmp:
        with pytest.raises(A.ErroPrestacao):
            A.desenhar({'slides': []}, {}, {}, tmp)


def test_rota_http():
    from fastapi.testclient import TestClient
    os.environ.setdefault('INTERNAL_API_SECRET', 'segredo-teste')
    from app import auth
    auth.INTERNAL_API_SECRET = auth.INTERNAL_API_SECRET or 'segredo-teste'
    from app.main import app
    cli = TestClient(app)
    h = {'X-Internal-Secret': auth.INTERNAL_API_SECRET}
    corpo = {'roteiro': {'slides': [{'tipo': 'capa', 'condominio': 'Cond', 'data': 'hoje'},
                                    {'tipo': 'fotos', 'titulo': 'F', 'fotos': [{'id': 'f1'}]}, {'tipo': 'encerramento'}]},
             'fotos': [{'id': 'f1', 'b64': base64.b64encode(_foto()).decode()}], 'documentos': [],
             'logo_cliente_b64': base64.b64encode(_logo()).decode()}
    r = cli.post('/v13/apresentacao', json=corpo, headers=h)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j['slides'] == 3 and base64.b64decode(j['pdf_b64'])[:4] == b'%PDF' and isinstance(j['avisos'], list)
    assert cli.post('/v13/apresentacao', json={'roteiro': {}}, headers=h).status_code == 422
    r = cli.post('/v13/apresentacao', json={**corpo, 'documentos': [{'id': 'd1', 'b64': base64.b64encode(b'oi').decode()}]}, headers=h)
    assert r.status_code == 422 and r.json()['detail']['erro'] == 'entrada_invalida'
    assert cli.post('/v13/apresentacao', json=corpo).status_code in (401, 503)
